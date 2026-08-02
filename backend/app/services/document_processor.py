"""Document text extraction and processing pipeline."""

import io
import json
import logging
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation

from app.config import get_settings
from app.services.ocr import get_ocr_service
from app.utils.text_cleaner import clean_text, is_text_sufficient

logger = logging.getLogger(__name__)
settings = get_settings()


@dataclass
class ExtractedPage:
    """A single page or section of extracted text."""

    page_number: int
    text: str
    section: str | None = None
    ocr_used: bool = False


@dataclass
class ExtractionResult:
    """Complete extraction result for a document."""

    pages: list[ExtractedPage] = field(default_factory=list)
    file_type: str = ""
    ocr_used: bool = False
    total_pages: int = 0

    def to_dict(self) -> dict[str, Any]:
        return {
            "pages": [asdict(page) for page in self.pages],
            "metadata": {
                "file_type": self.file_type,
                "ocr_used": self.ocr_used,
                "total_pages": self.total_pages,
            },
        }

    @property
    def full_text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text)


class DocumentProcessor:
    """Extract and clean text from supported document formats."""

    def __init__(self) -> None:
        self.ocr_service = get_ocr_service()

    def process(self, file_data: bytes, file_type: str) -> ExtractionResult:
        """Route to the appropriate extractor based on file type."""
        extractors = {
            "pdf": self._extract_pdf,
            "docx": self._extract_docx,
            "pptx": self._extract_pptx,
            "txt": self._extract_txt,
            "png": self._extract_image,
            "jpg": self._extract_image,
            "jpeg": self._extract_image,
            "tiff": self._extract_image,
            "bmp": self._extract_image,
        }

        extractor = extractors.get(file_type.lower())
        if not extractor:
            raise ValueError(f"Unsupported file type: {file_type}")

        result = extractor(file_data)
        result.file_type = file_type
        result.total_pages = len(result.pages)

        # Clean all extracted text
        for page in result.pages:
            page.text = clean_text(page.text)

        result.ocr_used = any(page.ocr_used for page in result.pages)
        return result

    def _extract_pdf(self, file_data: bytes) -> ExtractionResult:
        """Extract text from PDF using PyMuPDF, with OCR fallback for scanned pages."""
        pages: list[ExtractedPage] = []
        doc = fitz.open(stream=file_data, filetype="pdf")

        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            ocr_used = False

            if not is_text_sufficient(text, settings.min_text_chars_for_ocr):
                if self.ocr_service.is_available:
                    try:
                        pix = page.get_pixmap(dpi=200)
                        image_data = pix.tobytes("png")
                        text = self.ocr_service.extract_from_image(image_data)
                        ocr_used = True
                    except Exception as exc:
                        logger.warning("OCR failed for PDF page %d: %s", page_num + 1, exc)
                else:
                    logger.warning(
                        "Insufficient text on page %d and Tesseract not available",
                        page_num + 1,
                    )

            pages.append(
                ExtractedPage(
                    page_number=page_num + 1,
                    text=text,
                    section=None,
                    ocr_used=ocr_used,
                )
            )

        doc.close()
        return ExtractionResult(pages=pages)

    def _extract_docx(self, file_data: bytes) -> ExtractionResult:
        """Extract text from DOCX with section detection from headings."""
        doc = DocxDocument(io.BytesIO(file_data))
        pages: list[ExtractedPage] = []
        current_section: str | None = None
        current_text: list[str] = []
        page_number = 1

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            style_name = paragraph.style.name if paragraph.style else ""
            if style_name.startswith("Heading"):
                if current_text:
                    pages.append(
                        ExtractedPage(
                            page_number=page_number,
                            text="\n".join(current_text),
                            section=current_section,
                        )
                    )
                    page_number += 1
                    current_text = []
                current_section = text
            else:
                current_text.append(text)

        if current_text:
            pages.append(
                ExtractedPage(
                    page_number=page_number,
                    text="\n".join(current_text),
                    section=current_section,
                )
            )

        if not pages:
            full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            pages.append(ExtractedPage(page_number=1, text=full_text, section=None))

        return ExtractionResult(pages=pages)

    def _extract_pptx(self, file_data: bytes) -> ExtractionResult:
        """Extract text from PPTX slides."""
        prs = Presentation(io.BytesIO(file_data))
        pages: list[ExtractedPage] = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            texts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            texts.append(paragraph.text.strip())

            pages.append(
                ExtractedPage(
                    page_number=slide_num,
                    text="\n".join(texts),
                    section=f"Slide {slide_num}",
                )
            )

        return ExtractionResult(pages=pages)

    def _extract_txt(self, file_data: bytes) -> ExtractionResult:
        """Extract plain text file content."""
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                text = file_data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = file_data.decode("utf-8", errors="replace")

        return ExtractionResult(
            pages=[ExtractedPage(page_number=1, text=text, section=None)]
        )

    def _extract_image(self, file_data: bytes) -> ExtractionResult:
        """Extract text from image using OCR."""
        if not self.ocr_service.is_available:
            raise RuntimeError(
                "Tesseract OCR is not installed. Install Tesseract to process image files."
            )

        text = self.ocr_service.extract_from_image(file_data)
        return ExtractionResult(
            pages=[
                ExtractedPage(
                    page_number=1,
                    text=text,
                    section=None,
                    ocr_used=True,
                )
            ],
            ocr_used=True,
        )

    def save_extraction_result(
        self, result: ExtractionResult, storage, owner_id: int, document_id: int
    ) -> str:
        """Serialize extraction result to JSON and store it."""
        relative_path = storage.generate_path(
            owner_id=owner_id,
            filename=f"extracted_{document_id}.json",
            prefix="extracted",
        )
        json_data = json.dumps(result.to_dict(), ensure_ascii=False, indent=2)
        storage.save(json_data.encode("utf-8"), relative_path)
        return relative_path

    def load_extraction_result(self, storage, extracted_text_path: str) -> ExtractionResult:
        """Load a previously saved extraction result."""
        data = json.loads(storage.read(extracted_text_path).decode("utf-8"))
        pages = [
            ExtractedPage(
                page_number=p["page_number"],
                text=p["text"],
                section=p.get("section"),
                ocr_used=p.get("ocr_used", False),
            )
            for p in data.get("pages", [])
        ]
        metadata = data.get("metadata", {})
        return ExtractionResult(
            pages=pages,
            file_type=metadata.get("file_type", ""),
            ocr_used=metadata.get("ocr_used", False),
            total_pages=metadata.get("total_pages", len(pages)),
        )


def get_document_processor() -> DocumentProcessor:
    """Return document processor instance."""
    return DocumentProcessor()
